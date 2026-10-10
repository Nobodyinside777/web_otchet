"""Расчёт и сборка презентации.

Цифры проверяем на маленькой синтетической таблице, где суммы легко посчитать в уме.
Сборку PPTX — на обезличенном образце.
"""

import io
from datetime import date

import polars as pl
import pytest
from pptx import Presentation

from app.loader import load
from app.report import OTHERS, Plan, ReportParams, build_report, month_headers, report_months
from app.report_pptx import build_pptx, detail_pages, fmt_mln
from tests.conftest import SAMPLE

M = 1_000_000


def row(mvz, status, prob, revenue, margin, *, year=2026, month=10, contract=10, customer="А", manager="Иванов"):
    return {
        "year": year,
        "month": month,
        "mvz": mvz,
        "order_status": status,
        "probability": prob,
        "revenue": revenue * M,
        "margin_rub": margin * M,
        "contract_month": contract,
        "customer": customer,
        "description": "Описание",
        "contract_status_short": "Статус",
        "manager": manager,
    }


@pytest.fixture
def df():
    return pl.DataFrame(
        [
            row("ЦПА", "Поставлено", 1.0, 10, 2),  # признано
            row("ЦПА", "Законтрактовано", 1.0, 20, 4),  # законтрактовано, факт октября
            row("ЦПА", None, 1.0, 5, 1, contract=9),  # без статуса, 100 % — тоже законтрактовано
            row("ЦПА", "ТКП Отправлено", 0.5, 30, 3, month=11, contract=11, customer="Б"),  # высокая
            row("ЦИТО", "Заявка", 0.5, 40, 8, month=12, contract=10, customer="В"),  # высокая
            row("ЦИТО", "Заявка", 0.25, 7, 1),  # потенциал
            row("ЦИТО", "Поставлено", 1.0, 100, 50, year=2025),  # прошлый год — не в «Признано»
        ]
    )


@pytest.fixture
def params():
    return ReportParams(
        date(2026, 10, 10),
        plans={"ЦПА": Plan(50 * M, 10 * M, 60 * M, 12 * M), "ЦИТО": Plan(40 * M, 8 * M, 30 * M, 6 * M)},
    )


def test_months_and_labels(params):
    assert report_months(params) == [(2026, 10), (2026, 11), (2026, 12)]
    assert month_headers(params) == ["Октябрь Факт", "Октябрь", "Ноябрь", "Декабрь"]
    assert params.label_forecast() == "БП 9+3"
    assert params.label_budget() == "БП'26"
    assert report_months(ReportParams(date(2026, 11, 1))) == [(2026, 11), (2026, 12), (2027, 1)]


def test_company_kpi(df, params):
    report = build_report(df, params)
    revenue, margin, percent = report.summaries[0].kpi
    assert [round(v / M, 3) for v in revenue.values()] == [
        10, 25, 70, 105,  # признано, законтракт., высокая, итого
        90, 90, 15, 7,  # БП 9+3, БП'26, отклонение (105 − 90), потенциал
    ]  # fmt: skip
    assert margin.total == pytest.approx(2 * M + 5 * M + 11 * M)
    assert percent.total == pytest.approx(18 / 105)
    assert percent.values()[0] == pytest.approx(0.2)


def test_funnel_and_plan(df, params):
    company = build_report(df, params).summaries[0]
    funnel = {r.label: [v / M for v in r.values] for r in company.funnel}
    # Факт октября: контрактация в октябре и 100 % (в том числе уже поставленное). Месяцы — только 50 %.
    assert funnel["ЦПА"] == [10 + 20, 0, 30, 0]
    assert funnel["ЦИТО"] == [0, 40, 0, 0]
    assert funnel["Всего воронка продаж"] == [30, 40, 30, 0]
    plan = {r.label: [v / M for v in r.values] for r in company.plan}
    # План по выручке — по месяцу выручки: строка ЦИТО на 50 % идёт в декабре,
    # строка с контрактацией в сентябре, но выручкой в октябре — в факт октября.
    assert plan["Всего выручка"] == [10 + 20 + 5, 0, 30, 40]
    assert plan["Всего марж. прибыль"] == [2 + 4 + 1, 0, 3, 8]


def test_slides_per_mvz(df, params):
    report = build_report(df, params)
    titles = [s.title for s in report.summaries]
    assert titles == ["Исполнение БП'26  БЕ АСУТП", 'Исполнение БП\'26 по МВЗ "ЦИТО"', 'Исполнение БП\'26 по МВЗ "ЦПА"']
    cpa = report.summaries[2]
    assert cpa.kpi[0].budget == 60 * M
    # Детализация: только вероятность 50 %.
    assert [d.title for d in report.details] == [
        'Детализация контрактации по МВЗ "ЦИТО"',
        'Детализация контрактации по МВЗ "ЦПА"',
    ]
    assert [r[1] for r in report.details[1].rows] == ["Б"]


def test_filtered_rows_change_report(df, params):
    """Отчёт строится по тем строкам, что пришли: фильтр по МВЗ оставляет один слайд МВЗ."""
    report = build_report(df.filter(pl.col("mvz") == "ЦПА"), params)
    assert len(report.summaries) == 2
    assert report.summaries[0].kpi[0].budget == 60 * M


def test_customers_top_and_others(params):
    rows = [row("ЦПА", "Заявка", 0.5, i + 1, 0, customer=f"К{i}") for i in range(10)]
    report = build_report(pl.DataFrame(rows), params, customers=3)
    labels = [r.label for r in report.summaries[1].funnel]
    assert labels == ["К9", "К8", "К7", OTHERS, "Всего контрактация"]
    assert report.summaries[1].funnel[3].values[1] == pytest.approx(sum(range(1, 8)) * M)


def test_number_format():
    assert fmt_mln(12_345_678) == "12,3"
    assert fmt_mln(-1_234_567_890) == "-1 234,6"
    assert fmt_mln(10_000) == "0"
    assert fmt_mln(10_000, blank_zero=True) == ""


def test_long_detail_is_split_into_pages():
    rows = [[10, "Заказчик", "Описание " * 10, 1.0, "Статус " * 20, "Иванов"] for _ in range(60)]
    pages = detail_pages(rows)
    assert len(pages) > 1
    assert sum(len(p) for p in pages) == 60


def test_pptx_from_sample():
    df = load(SAMPLE).df
    params = ReportParams(date(2026, 10, 10), title_texts=("Заголовок", "Подзаголовок", "Подпись"))
    report = build_report(df, params)
    prs = Presentation(io.BytesIO(build_pptx(report)))

    mvz_count = df.get_column("mvz").n_unique()
    assert len(prs.slides) >= 1 + 1 + mvz_count + mvz_count  # титул, общий, по МВЗ, детализации
    first_table = next(sh.table for sh in prs.slides[0].shapes if sh.has_table)
    assert [c.text for c in first_table.rows[0].cells] == ["10", "10", "2026"]
    texts = [sh.text_frame.text for sh in sorted(prs.slides[0].shapes, key=lambda s: s.top)
             if sh.has_text_frame and sh.text_frame.text]
    assert texts == ["Заголовок", "Подзаголовок", "Подпись"]
    assert prs.slides[1].shapes.title.text == "Исполнение БП'26  БЕ АСУТП"


def test_template_has_no_texts_or_names():
    """В шаблоне в git нет текстов титула, ФИО и названий: только оформление."""
    import re
    import zipfile

    from app.report_pptx import TEMPLATE

    z = zipfile.ZipFile(TEMPLATE)
    slide = z.read("ppt/slides/slide1.xml").decode("utf-8")
    assert not [t for t in re.findall(r"<a:t>([^<]*)</a:t>", slide) if t.strip()]
    core = z.read("docProps/core.xml").decode("utf-8")
    assert "<dc:creator></dc:creator>" in core or "<dc:creator/>" in core
    assert "TitlesOfParts" not in z.read("docProps/app.xml").decode("utf-8")
    assert not [n for n in z.namelist() if "author" in n.lower() or "comment" in n.lower()]
