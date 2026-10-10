"""Сборка презентации PPTX из посчитанного отчёта (app/report.py).

За основу берётся шаблон templates/report_template.pptx: в нём фирменные макеты
и титульный слайд. Слайды с таблицами строятся заново по образцу присланной презентации:
шрифт PT Root UI, серая шапка, оранжевая строка итога, чёрная строка маржи.

PowerPoint для этого не нужен: файл собирает библиотека python-pptx.
"""

import io
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

from app.report import DETAIL_HEADERS, DETAIL_REVENUE, DetailSlide, KpiRow, MonthRow, Report, SummarySlide

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "report_template.pptx"

FONT = "PT Root UI"
TITLE_FONT = "PT Root UI Bold"
BLACK = "000000"
WHITE = "FFFFFF"
HEADER_FILL = "F5F5F5"
BORDER = "A6A6A6"
ORANGE = "FE5F19"
DARK = "262626"
DARK_BORDER = "BFBFBF"
DETAIL_HEADER_FILL = "F2F2F2"

# Макеты второго мастера шаблона (по названию, как в PowerPoint).
LAYOUT_ONE_LINE = "Белый со стр и заголовком"
LAYOUT_TWO_LINES = "1_Белый со стр и заголовком в 2 строки"

# Геометрия как в образце (EMU: 914400 на дюйм).
TABLE_X = 493714
TABLE_Y = 1086928
TABLE_W = 11013917
TABLE_BOTTOM = 6400000  # ниже — номер слайда
SUMMARY_COLS = [1891959, 985909, 985909, 1191690, 1191690, 1191690, 1191690, 1191690, 1191690]
SUMMARY_ROW_H = 268590
SUMMARY_ROW_MIN = 200000
SUMMARY_HEADER_H = 537180
SUMMARY_FONT = 10

DETAIL_COLS = [548165, 909592, 3054062, 642537, 4891318, 931679]
DETAIL_X = 519023
DETAIL_Y = 1150224
DETAIL_HEADER_H = 499543
DETAIL_LINE_H = 162163  # высота строки в одну строчку текста
DETAIL_FONT = 9
CHAR_W = 9 * 12700 * 0.55  # средняя ширина символа шрифта 9 pt, EMU (для оценки переносов)

LINE_X, LINE_Y, LINE_H = 493714, 317628, 618997  # вертикальная черта слева от заголовка


# ---------- форматирование чисел ----------


def fmt_mln(value: float, blank_zero: bool = False) -> str:
    """Рубли -> миллионы с одним знаком: 12345678 -> '12,3'. Отрицательные со знаком минус."""
    mln = value / 1_000_000
    if blank_zero and abs(mln) < 0.05:
        return ""
    text = f"{mln:,.1f}".replace(",", " ").replace(".", ",")
    return "0" if text in ("0,0", "-0,0") else text


def fmt_pct(value: float) -> str:
    return f"{round(value * 100):.0f}%"


def fmt_month(value) -> str:
    return str(value) if value != "" else ""


# ---------- низкоуровневые помощники для ячеек ----------


def _line(tag: str, color: str | None, width: int = 6350):
    """Граница ячейки (<a:lnL> и т. п.); color=None — без линии."""
    ln = etree.SubElement(etree.Element("dummy"), qn(f"a:{tag}"))
    if color is None:
        etree.SubElement(ln, qn("a:noFill"))
        return ln
    ln.set("w", str(width))
    fill = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
    return ln


def style_cell(cell, text: str, *, size: int, bold=False, color=BLACK, fill=None,
               align=PP_ALIGN.CENTER, borders=(BORDER, BORDER, BORDER, BORDER), border_w=6350):  # fmt: skip
    """Текст, шрифт, заливка и границы ячейки. borders — (слева, справа, сверху, снизу)."""
    cell.margin_left = cell.margin_right = Emu(8802)
    cell.margin_top = Emu(8802)
    cell.margin_bottom = Emu(0)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    tf = cell.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.name = FONT
    font.color.rgb = RGBColor.from_string(color)

    # Границы и заливку пишем в XML сами: python-pptx умеет только заливку.
    # Порядок важен для PowerPoint: lnL, lnR, lnT, lnB, затем заливка.
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        tc_pr.remove(child)
    for tag, color_ in zip(("lnL", "lnR", "lnT", "lnB"), borders):
        tc_pr.append(_line(tag, color_, border_w))
    if fill:
        solid = etree.SubElement(tc_pr, qn("a:solidFill"))
        etree.SubElement(solid, qn("a:srgbClr")).set("val", fill)
    else:
        etree.SubElement(tc_pr, qn("a:noFill"))


def new_table(slide, rows: int, widths: list[int], x: int, y: int, heights: list[int]):
    shape = slide.shapes.add_table(rows, len(widths), Emu(x), Emu(y), Emu(sum(widths)), Emu(sum(heights)))
    table = shape.table
    # Убираем встроенный стиль PowerPoint (полосы, синяя шапка): оформляем ячейки сами.
    table.first_row = False
    table.horz_banding = False
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    for style in tbl_pr.findall(qn("a:tableStyleId")):
        tbl_pr.remove(style)
    for i, w in enumerate(widths):
        table.columns[i].width = Emu(w)
    for i, h in enumerate(heights):
        table.rows[i].height = Emu(h)
    return table


def add_slide(prs, layout_name: str, title: str):
    layout = next(lo for master in prs.slide_masters for lo in master.slide_layouts if lo.name == layout_name)
    slide = prs.slides.add_slide(layout)
    title_shape = slide.shapes.title
    run = title_shape.text_frame.paragraphs[0].add_run()
    run.text = title
    run.font.name = TITLE_FONT
    line = slide.shapes.add_connector(1, Emu(LINE_X), Emu(LINE_Y), Emu(LINE_X), Emu(LINE_Y + LINE_H))
    line.line.width = Emu(19050)
    line.line.color.rgb = RGBColor.from_string("2B2B2A")
    return slide


# ---------- титульный слайд ----------


def fill_title_slide(prs, report: Report) -> None:
    """Титул: три надписи (заголовок, подзаголовок, подпись) и ячейки день | месяц | год отчёта."""
    slide = prs.slides[0]
    d = report.params.report_date
    table = next(sh.table for sh in slide.shapes if sh.has_table)
    for cell, text in zip(table.rows[0].cells, [f"{d.day:02d}", f"{d.month:02d}", str(d.year)]):
        _set_text(cell.text_frame, text)

    # Надписи в шаблоне пустые; берём их сверху вниз.
    boxes = sorted(
        (sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.paragraphs[0].runs),
        key=lambda sh: sh.top,
    )
    for box, text in zip(boxes, report.params.title_texts):
        _set_text(box.text_frame, text)


def _set_text(text_frame, text: str) -> None:
    """Меняет только текст первого фрагмента, оформление шаблона остаётся."""
    runs = text_frame.paragraphs[0].runs
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


# ---------- сводный слайд ----------


def summary_slide(prs, report: Report, data: SummarySlide) -> None:
    params = report.params
    kpi_headers = [
        "Показатель, млн. руб.", "Признано", "Законтракт.", "Высокая вероятность", "Итого",
        params.label_forecast(), params.label_budget(), "Отклонение", "Потенциал",
    ]  # fmt: skip
    month_headers = [*report.month_headers, "Итого"]

    # Порядок строк: шапка, показатели, пусто, шапка воронки, воронка, пусто, шапка плана, план.
    layout: list[tuple[str, object]] = [("kpi_header", None)]
    layout += [("kpi", r) for r in data.kpi]
    layout += [("gap", None), ("month_header", "Заказчик / Месяц контрактации")]
    layout += [("month", r) for r in data.funnel]
    layout += [("gap", None), ("month_header", "План по выручке")]
    layout += [("month", r) for r in data.plan]

    headers = sum(1 for kind, _ in layout if kind == "kpi_header")
    body = len(layout) - headers
    space = TABLE_BOTTOM - TABLE_Y - SUMMARY_HEADER_H
    row_h = max(SUMMARY_ROW_MIN, min(SUMMARY_ROW_H, space // body))
    heights = [SUMMARY_HEADER_H if kind == "kpi_header" else row_h for kind, _ in layout]

    two_lines = len(data.title) > 55
    slide = add_slide(prs, LAYOUT_TWO_LINES if two_lines else LAYOUT_ONE_LINE, data.title)
    table = new_table(slide, len(layout), SUMMARY_COLS, TABLE_X, TABLE_Y, heights)

    for i, (kind, item) in enumerate(layout):
        cells = list(table.rows[i].cells)
        if kind == "kpi_header":
            for cell, text in zip(cells, kpi_headers):
                style_cell(cell, text, size=SUMMARY_FONT, bold=True, fill=HEADER_FILL)
        elif kind == "kpi":
            _kpi_row(cells, item)
        elif kind == "gap":
            for cell in cells:
                style_cell(cell, "", size=SUMMARY_FONT, borders=(None, None, None, None))
        elif kind == "month_header":
            _month_row(table, i, item, month_headers, bold=True, fill=HEADER_FILL)
        else:
            _month_data_row(table, i, item)


def _kpi_row(cells, row: KpiRow) -> None:
    style_cell(cells[0], row.label, size=SUMMARY_FONT, bold=True, align=PP_ALIGN.LEFT)
    fmt = fmt_pct if row.percent else fmt_mln
    for cell, value in zip(cells[1:], row.values()):
        style_cell(cell, fmt(value), size=SUMMARY_FONT, bold=True)


def _month_row(table, i: int, label: str, values: list[str], *, bold=False, fill=None,
               color=BLACK, borders=(BORDER,) * 4) -> None:  # fmt: skip
    """Строка «подпись на 4 колонки + 5 значений»."""
    first = table.cell(i, 0)
    first.merge(table.cell(i, 3))
    style_cell(first, label, size=SUMMARY_FONT, bold=bold, fill=fill, color=color, borders=borders)
    for col, text in enumerate(values, start=4):
        style_cell(table.cell(i, col), text, size=SUMMARY_FONT, bold=bold, fill=fill, color=color, borders=borders)


def _month_data_row(table, i: int, row: MonthRow) -> None:
    if row.level == "total":
        values = [fmt_mln(v) for v in [*row.values, row.total]]
        _month_row(table, i, row.label, values, fill=ORANGE, color=WHITE)
    elif row.level == "margin":
        values = [fmt_mln(v) for v in [*row.values, row.total]]
        _month_row(table, i, row.label, values, bold=True, fill=DARK, color=WHITE, borders=(DARK_BORDER,) * 4)
    else:
        values = [fmt_mln(v, blank_zero=True) for v in [*row.values, row.total]]
        _month_row(table, i, row.label, values)


# ---------- детализация ----------


def _lines(text: str, width: int) -> int:
    """Сколько строчек займёт текст в колонке шириной width (грубая оценка)."""
    usable = max(width - 2 * 8802, 1)
    return max(1, math.ceil(len(str(text)) * CHAR_W / usable))


def detail_pages(rows: list[list]) -> list[list[list]]:
    """Делит строки детализации на страницы так, чтобы таблица помещалась на слайд."""
    space = TABLE_BOTTOM - DETAIL_Y - DETAIL_HEADER_H
    pages, page, used = [], [], 0
    for row in rows:
        texts = [fmt_month(row[0]), row[1], row[2], "0,0", row[4], row[5]]
        h = DETAIL_LINE_H * max(_lines(t, w) for t, w in zip(texts, DETAIL_COLS))
        if page and used + h > space:
            pages.append(page)
            page, used = [], 0
        page.append(row)
        used += h
    pages.append(page)  # даже пустая детализация даёт один слайд с шапкой
    return pages


def detail_slides(prs, data: DetailSlide) -> None:
    pages = detail_pages(data.rows)
    for n, page in enumerate(pages):
        title = data.title if n == 0 else f"{data.title} (продолжение)"
        slide = add_slide(prs, LAYOUT_TWO_LINES, title)
        heights = [DETAIL_HEADER_H] + [DETAIL_LINE_H] * max(len(page), 1)
        table = new_table(slide, len(heights), DETAIL_COLS, DETAIL_X, DETAIL_Y, heights)
        for col, text in enumerate(DETAIL_HEADERS):
            style_cell(table.cell(0, col), text, size=DETAIL_FONT, bold=True, fill=DETAIL_HEADER_FILL,
                       borders=(BLACK,) * 4)  # fmt: skip
        if not page:
            table.cell(1, 0).merge(table.cell(1, len(DETAIL_COLS) - 1))
            style_cell(table.cell(1, 0), "Нет сделок с вероятностью 50 %", size=DETAIL_FONT, borders=(BLACK,) * 4)
            continue
        for i, row in enumerate(page, start=1):
            texts = [fmt_month(row[0]), row[1], row[2], fmt_mln(row[DETAIL_REVENUE]), row[4], row[5]]
            for col, text in enumerate(texts):
                align = PP_ALIGN.LEFT if col == 4 else PP_ALIGN.CENTER
                style_cell(table.cell(i, col), str(text), size=DETAIL_FONT, align=align, borders=(BLACK,) * 4)


# ---------- вся презентация ----------


def build_pptx(report: Report, template: Path = TEMPLATE) -> bytes:
    """Готовая презентация в виде байтов (для кнопки «Скачать»)."""
    prs = Presentation(str(template))
    fill_title_slide(prs, report)
    for summary in report.summaries:
        summary_slide(prs, report, summary)
    for detail in report.details:
        detail_slides(prs, detail)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
